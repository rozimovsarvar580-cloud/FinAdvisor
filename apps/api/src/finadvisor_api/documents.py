import json
from decimal import Decimal, InvalidOperation
from importlib.resources import files
from io import BytesIO
from typing import Annotated, Literal
from xml.sax.saxutils import escape

from fastapi import APIRouter, Depends, HTTPException, Path, status
from fastapi.responses import Response
from finance_engine.operations import CapexItem, StaffPosition
from finance_engine.restaurant import (
    RestaurantProjection,
    calculate_restaurant_projection,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import model_validator
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, StyleSheet1, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from finadvisor_api.auth import get_current_user
from finadvisor_api.calc_schemas import StrictInput
from finadvisor_api.models import User
from finadvisor_api.plan_limits import require_plan_feature
from finadvisor_api.plan_schemas import PlanGenerateRequest, PlanGenerateResponse, PlanSection
from finadvisor_api.readiness import calculate_bank_readiness

router = APIRouter(prefix="/documents", tags=["documents"])
MONTHLY_PROJECTION_MONTHS = 12
MONEY_NUMBER_FORMAT = "#,##0.00"

DOCUMENT_MESSAGES = json.loads(
    files("finadvisor_api").joinpath("document_messages.json").read_text(encoding="utf-8")
)
FINANCIAL_KEYS = (
    "monthly_revenue",
    "monthly_payroll",
    "startup_capex",
    "total_startup_cost",
    "monthly_fixed_costs",
    "monthly_variable_cost",
    "monthly_operating_profit",
    "annual_operating_profit",
    "break_even_revenue",
    "target_monthly_revenue",
    "loan_first_payment",
    "loan_total_payment",
    "loan_total_interest",
    "loan_term_months",
    "annual_debt_service",
    "dscr",
)
CURRENCY_KEYS = frozenset(
    {
        "monthly_revenue",
        "monthly_payroll",
        "startup_capex",
        "total_startup_cost",
        "monthly_fixed_costs",
        "monthly_variable_cost",
        "monthly_operating_profit",
        "annual_operating_profit",
        "break_even_revenue",
        "target_monthly_revenue",
        "loan_first_payment",
        "loan_total_payment",
        "loan_total_interest",
        "annual_debt_service",
    }
)


class PlanDocumentRequest(StrictInput):
    locale: Literal["uz", "ru", "en"]
    plan: PlanGenerateResponse

    @model_validator(mode="after")
    def validate_financial_values(self) -> "PlanDocumentRequest":
        for key in FINANCIAL_KEYS:
            value = getattr(self.plan.calculations, key)
            if value is None:
                continue
            try:
                amount = Decimal(str(value))
            except InvalidOperation as error:
                raise ValueError(f"{key} must be a decimal value") from error
            if not amount.is_finite():
                raise ValueError(f"{key} must be finite")
        return self


def _messages(locale: str) -> dict[str, object]:
    return DOCUMENT_MESSAGES[locale]


def _excel_text(value: str) -> str:
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r")):
        return f"'{value}"
    return value


def _restaurant_projection(
    assumptions: PlanGenerateRequest,
) -> RestaurantProjection:
    try:
        return calculate_restaurant_projection(
            average_check=assumptions.average_check,
            daily_customers=assumptions.daily_customers,
            operating_days_per_month=assumptions.operating_days_per_month,
            monthly_rent=assumptions.monthly_rent,
            monthly_utilities=assumptions.monthly_utilities,
            variable_cost_ratio_percent=assumptions.variable_cost_ratio_percent,
            staff=[
                StaffPosition(item.role, item.employee_count, item.monthly_salary)
                for item in assumptions.staff
            ],
            equipment=[
                CapexItem(item.name, item.quantity, item.unit_cost)
                for item in assumptions.equipment
            ],
            other_startup_cost=assumptions.other_startup_cost,
            loan_amount=assumptions.loan_amount,
            annual_interest_rate_percent=assumptions.annual_interest_rate_percent,
            loan_term_months=assumptions.loan_term_months,
            target_monthly_profit=assumptions.target_monthly_profit,
        )
    except ValueError as error:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="documents.invalid_plan_assumptions",
        ) from error


def _projection_values(
    projection: RestaurantProjection,
) -> dict[str, Decimal | int | None]:
    loan = projection.loan_schedule
    return {
        "monthly_revenue": projection.monthly_revenue,
        "monthly_payroll": projection.monthly_payroll.monthly_payroll,
        "startup_capex": projection.startup_capex.total_cost,
        "total_startup_cost": projection.total_startup_cost,
        "monthly_fixed_costs": projection.monthly_fixed_costs,
        "monthly_variable_cost": projection.monthly_variable_cost,
        "monthly_operating_profit": projection.monthly_operating_profit,
        "annual_operating_profit": projection.annual_operating_profit,
        "break_even_revenue": projection.break_even.break_even_revenue,
        "target_monthly_revenue": projection.target_revenue.required_revenue,
        "loan_first_payment": loan.monthly_payment if loan is not None else None,
        "loan_total_payment": loan.total_payment if loan is not None else None,
        "loan_total_interest": loan.total_interest if loan is not None else None,
        "loan_term_months": loan.term_months if loan is not None else None,
        "annual_debt_service": projection.annual_debt_service,
        "dscr": projection.dscr.dscr if projection.dscr is not None else None,
    }


def _style_header(
    sheet: Worksheet,
    row: int,
    fill: PatternFill,
    font: Font,
) -> None:
    for cell in sheet[row]:
        cell.fill = fill
        cell.font = font


def _pdf_table(rows: list[list[str]], col_widths: list[float]) -> Table:
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "FinAdvisor"),
                ("FONTNAME", (0, 0), (-1, 0), "FinAdvisor-Bold"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F6F8")),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#183B56")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D8E0E5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _auxiliary_pdf_setup(
    title: str,
) -> tuple[BytesIO, SimpleDocTemplate, StyleSheet1]:
    font_path = files("finadvisor_api").joinpath("fonts", "NotoSans.ttf")
    pdfmetrics.registerFont(TTFont("FinAdvisor", str(font_path)))
    pdfmetrics.registerFont(TTFont("FinAdvisor-Bold", str(font_path)))
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "FinAdvisorAuxTitle",
            parent=styles["Title"],
            fontName="FinAdvisor-Bold",
            alignment=TA_CENTER,
            textColor=colors.HexColor("#183B56"),
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            "FinAdvisorAuxHeading",
            parent=styles["Heading2"],
            fontName="FinAdvisor-Bold",
            textColor=colors.HexColor("#183B56"),
            spaceBefore=10,
            spaceAfter=7,
        )
    )
    styles.add(
        ParagraphStyle(
            "FinAdvisorAuxBody",
            parent=styles["BodyText"],
            fontName="FinAdvisor",
            leading=15,
            spaceAfter=8,
        )
    )
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title,
    )
    return output, document, styles


def _matching_sections(
    request: PlanDocumentRequest,
    aliases: tuple[str, ...],
) -> list[PlanSection]:
    return [
        section
        for section in request.plan.sections
        if any(alias in section.title.casefold() for alias in aliases)
    ]


def render_bank_pack_pdf(request: PlanDocumentRequest) -> bytes:
    labels = _messages(request.locale)
    assumptions = request.plan.assumptions
    projection = _restaurant_projection(assumptions)
    readiness_text = " ".join(
        [request.plan.summary]
        + [f"{section.title} {section.content}" for section in request.plan.sections]
    )
    readiness = calculate_bank_readiness(readiness_text)
    output, document, styles = _auxiliary_pdf_setup(str(labels["bank_pack_title"]))
    story = [
        Paragraph(
            escape(str(labels["bank_pack_title"])),
            styles["FinAdvisorAuxTitle"],
        ),
        Paragraph(escape(assumptions.business_name), styles["FinAdvisorAuxTitle"]),
        Paragraph(escape(str(labels["loan_request"])), styles["FinAdvisorAuxHeading"]),
        _pdf_table(
            [
                [str(labels["table_metric"]), str(labels["table_value"])],
                [
                    str(labels["requested_amount"]),
                    f"{assumptions.loan_amount:,.2f} {labels['currency']}",
                ],
                [
                    str(labels["annual_rate"]),
                    f"{assumptions.annual_interest_rate_percent}%"
                    if assumptions.loan_amount
                    else str(labels["status"]["no_loan"]),
                ],
                [
                    str(labels["loan_term"]),
                    str(assumptions.loan_term_months)
                    if assumptions.loan_amount
                    else str(labels["status"]["no_loan"]),
                ],
            ],
            [85 * mm, 75 * mm],
        ),
        Paragraph(escape(str(labels["repayment_schedule"])), styles["FinAdvisorAuxHeading"]),
    ]
    loan_schedule = projection.loan_schedule
    if loan_schedule is None:
        story.append(
            Paragraph(
                escape(str(labels["status"]["no_loan"])),
                styles["FinAdvisorAuxBody"],
            )
        )
    else:
        loan_rows = [
            [
                str(labels["columns"]["month"]),
                str(labels["columns"]["payment"]),
                str(labels["columns"]["principal"]),
                str(labels["columns"]["interest"]),
                str(labels["columns"]["remaining_balance"]),
            ]
        ]
        for payment in loan_schedule.payments:
            loan_rows.append(
                [
                    str(payment.month),
                    f"{payment.payment:,.2f}",
                    f"{payment.principal:,.2f}",
                    f"{payment.interest:,.2f}",
                    f"{payment.remaining_balance:,.2f}",
                ]
            )
        story.append(_pdf_table(loan_rows, [20 * mm, 35 * mm, 35 * mm, 35 * mm, 35 * mm]))

    story.extend(
        [
            Paragraph(escape(str(labels["collateral"])), styles["FinAdvisorAuxHeading"]),
            Paragraph(
                escape(str(labels["collateral_not_supplied"])),
                styles["FinAdvisorAuxBody"],
            ),
            Paragraph(
                escape(str(labels["bank_readiness"])),
                styles["FinAdvisorAuxHeading"],
            ),
            Paragraph(
                f"{readiness.score}/100",
                styles["FinAdvisorAuxTitle"],
            ),
            Paragraph(
                escape(str(labels["readiness_disclaimer"])),
                styles["FinAdvisorAuxBody"],
            ),
            Paragraph(
                f"{escape(str(labels['checklist_version']))}: 1",
                styles["FinAdvisorAuxBody"],
            ),
        ]
    )
    readiness_rows = [
        [
            str(labels["criterion"]),
            str(labels["evidence_score"]),
            str(labels["evidence_status"]),
        ]
    ]
    for criterion in readiness.criteria:
        readiness_rows.append(
            [
                str(labels["readiness_criteria"][criterion.key]),
                f"{criterion.score}/{criterion.max_score}",
                str(
                    labels["evidence_matched"]
                    if criterion.present
                    else labels["evidence_not_matched"]
                ),
            ]
        )
    story.append(_pdf_table(readiness_rows, [80 * mm, 35 * mm, 45 * mm]))
    story.append(Paragraph(escape(str(labels["disclaimer"])), styles["FinAdvisorAuxBody"]))
    document.build(story)
    return output.getvalue()


def render_synergy_pdf(request: PlanDocumentRequest) -> bytes:
    labels = _messages(request.locale)
    assumptions = request.plan.assumptions
    projection = _restaurant_projection(assumptions)
    output, document, styles = _auxiliary_pdf_setup(str(labels["synergy_title"]))
    story = [
        Paragraph(escape(str(labels["synergy_title"])), styles["FinAdvisorAuxTitle"]),
        Paragraph(escape(assumptions.business_name), styles["FinAdvisorAuxTitle"]),
        Paragraph(
            escape(str(labels["strategy_note"])),
            styles["FinAdvisorAuxBody"],
        ),
        Paragraph(
            escape(str(labels["operating_concept"])),
            styles["FinAdvisorAuxHeading"],
        ),
        Paragraph(escape(assumptions.menu_summary), styles["FinAdvisorAuxBody"]),
    ]

    strategy_sections = (
        ("operating_plan", ("operation", "операц", "faoliyat")),
        ("team_plan", ("staff", "персонал", "xodim", "ishchi")),
        ("risk_management", ("risk", "риск", "xavf", "tavakkal")),
    )
    for label_key, aliases in strategy_sections:
        story.append(Paragraph(escape(str(labels[label_key])), styles["FinAdvisorAuxHeading"]))
        sections = _matching_sections(request, aliases)
        if sections:
            for section in sections:
                story.append(
                    Paragraph(
                        escape(section.content).replace("\n", "<br/>"),
                        styles["FinAdvisorAuxBody"],
                    )
                )
        else:
            story.append(
                Paragraph(
                    escape(str(labels["information_not_supplied"])),
                    styles["FinAdvisorAuxBody"],
                )
            )

    story.append(
        Paragraph(
            escape(str(labels["operating_resources"])),
            styles["FinAdvisorAuxHeading"],
        )
    )
    resources = [
        [str(labels["table_metric"]), str(labels["table_value"])],
        [str(labels["cover_location"]), escape(assumptions.location)],
        [
            str(labels["cover_format"]),
            str(labels["formats"][assumptions.restaurant_format]),
        ],
        [str(labels["assumptions"]["seats"]), str(assumptions.seats)],
        [
            str(labels["assumptions"]["operating_days_per_month"]),
            str(assumptions.operating_days_per_month),
        ],
    ]
    story.append(_pdf_table(resources, [85 * mm, 75 * mm]))

    story.append(
        Paragraph(
            escape(str(labels["operating_financials"])),
            styles["FinAdvisorAuxHeading"],
        )
    )
    financial_values = _projection_values(projection)
    financial_rows = [
        [str(labels["table_metric"]), str(labels["table_value"])],
    ]
    for key in (
        "monthly_revenue",
        "monthly_payroll",
        "monthly_fixed_costs",
        "monthly_variable_cost",
        "monthly_operating_profit",
    ):
        value = financial_values[key]
        financial_rows.append(
            [
                str(labels["metrics"][key]),
                f"{value:,.2f} {labels['currency']}",
            ]
        )
    story.append(_pdf_table(financial_rows, [85 * mm, 75 * mm]))
    story.append(Paragraph(escape(str(labels["disclaimer"])), styles["FinAdvisorAuxBody"]))
    document.build(story)
    return output.getvalue()


def render_plan_pdf(request: PlanDocumentRequest) -> bytes:
    labels = _messages(request.locale)
    assumptions = request.plan.assumptions
    projection = _restaurant_projection(assumptions)
    financial_values = _projection_values(projection)
    font_path = files("finadvisor_api").joinpath("fonts", "NotoSans.ttf")
    pdfmetrics.registerFont(TTFont("FinAdvisor", str(font_path)))
    pdfmetrics.registerFont(TTFont("FinAdvisor-Bold", str(font_path)))

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            "FinAdvisorTitle",
            parent=styles["Title"],
            fontName="FinAdvisor-Bold",
            alignment=TA_CENTER,
            textColor=colors.HexColor("#183B56"),
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            "FinAdvisorBody",
            parent=styles["BodyText"],
            fontName="FinAdvisor",
            leading=15,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            "FinAdvisorHeading",
            parent=styles["Heading2"],
            fontName="FinAdvisor-Bold",
            textColor=colors.HexColor("#183B56"),
            spaceBefore=10,
            spaceAfter=7,
        )
    )
    styles.add(
        ParagraphStyle(
            "FinAdvisorCover",
            parent=styles["Heading2"],
            fontName="FinAdvisor",
            alignment=TA_CENTER,
            leading=18,
            spaceAfter=8,
        )
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=str(labels["title"]),
    )
    formats = labels["formats"]
    story = [
        Paragraph(escape(str(labels["title"])), styles["FinAdvisorTitle"]),
        Spacer(1, 30 * mm),
        Paragraph(escape(assumptions.business_name), styles["FinAdvisorTitle"]),
        Paragraph(
            escape(str(labels["cover_location"])) + ": " + escape(assumptions.location),
            styles["FinAdvisorCover"],
        ),
        Paragraph(
            escape(str(labels["cover_format"]))
            + ": "
            + escape(str(formats[assumptions.restaurant_format])),
            styles["FinAdvisorCover"],
        ),
        Paragraph(
            escape(str(labels["assumptions"]["seats"])) + ": " + str(assumptions.seats),
            styles["FinAdvisorCover"],
        ),
        Paragraph(
            escape(str(labels["assumptions"]["menu_summary"]))
            + ": "
            + escape(assumptions.menu_summary),
            styles["FinAdvisorBody"],
        ),
        PageBreak(),
        Paragraph(escape(str(labels["summary"])), styles["FinAdvisorHeading"]),
        Paragraph(escape(request.plan.summary).replace("\n", "<br/>"), styles["FinAdvisorBody"]),
    ]

    section_aliases = {
        "market": ("market", "рынок", "bozor"),
        "operations": ("operation", "операц", "faoliyat"),
        "staffing": ("staff", "персонал", "xodim", "ishchi"),
        "risks": ("risk", "риск", "xavf", "tavakkal"),
        "financials": ("financial", "финанс", "moliyaviy", "moliya"),
    }
    categorized: dict[str, list[PlanSection]] = {key: [] for key in section_aliases}
    additional_sections = []
    for section in request.plan.sections:
        title = section.title.casefold()
        category = next(
            (
                key
                for key, aliases in section_aliases.items()
                if any(alias in title for alias in aliases)
            ),
            None,
        )
        if category is None:
            additional_sections.append(section)
        else:
            categorized[category].append(section)

    for category in ("market", "operations", "staffing", "risks"):
        story.append(Paragraph(escape(str(labels[category])), styles["FinAdvisorHeading"]))
        sections = categorized[category]
        if sections:
            for section in sections:
                story.append(
                    Paragraph(
                        escape(section.content).replace("\n", "<br/>"),
                        styles["FinAdvisorBody"],
                    )
                )
        else:
            story.append(
                Paragraph(
                    escape(str(labels["information_not_supplied"])),
                    styles["FinAdvisorBody"],
                )
            )

    story.append(Paragraph(escape(str(labels["financials"])), styles["FinAdvisorHeading"]))
    for section in categorized["financials"]:
        story.extend(
            [
                Paragraph(escape(section.title), styles["FinAdvisorHeading"]),
                Paragraph(
                    escape(section.content).replace("\n", "<br/>"),
                    styles["FinAdvisorBody"],
                ),
            ]
        )
    financial_rows = [[str(labels["table_metric"]), str(labels["table_value"])]]
    for key in FINANCIAL_KEYS:
        value = financial_values[key]
        if value is None:
            continue
        display_value = f"{value:,.2f} {labels['currency']}" if key in CURRENCY_KEYS else str(value)
        financial_rows.append([str(labels["metrics"][key]), display_value])
    story.append(_pdf_table(financial_rows, [105 * mm, 55 * mm]))

    chart_keys = (
        "monthly_revenue",
        "monthly_payroll",
        "monthly_fixed_costs",
        "monthly_variable_cost",
        "monthly_operating_profit",
    )
    chart_values = [float(Decimal(str(financial_values[key]))) for key in chart_keys]
    chart = VerticalBarChart()
    chart.x = 28
    chart.y = 24
    chart.width = 445
    chart.height = 150
    chart.data = [chart_values]
    chart.categoryAxis.categoryNames = [str(labels["metrics"][key]) for key in chart_keys]
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.boxAnchor = "ne"
    chart.categoryAxis.labels.dy = -8
    chart.categoryAxis.labels.fontName = "FinAdvisor"
    chart.valueAxis.labels.fontName = "FinAdvisor"
    chart.valueAxis.valueMin = min(0.0, min(chart_values))
    chart.valueAxis.valueMax = max(1.0, max(chart_values))
    chart.bars[0].fillColor = colors.HexColor("#F2A541")
    chart.barWidth = 32
    drawing = Drawing(500, 205)
    drawing.add(chart)
    story.extend(
        [
            Paragraph(
                escape(str(labels["financial_chart"])),
                styles["FinAdvisorHeading"],
            ),
            drawing,
        ]
    )

    if additional_sections:
        story.append(
            Paragraph(
                escape(str(labels["additional_plan_details"])),
                styles["FinAdvisorHeading"],
            )
        )
        for section in additional_sections:
            story.append(Paragraph(escape(section.title), styles["FinAdvisorHeading"]))
            story.append(
                Paragraph(
                    escape(section.content).replace("\n", "<br/>"),
                    styles["FinAdvisorBody"],
                )
            )

    story.extend(
        [
            PageBreak(),
            Paragraph(escape(str(labels["appendix"])), styles["FinAdvisorHeading"]),
        ]
    )
    assumption_rows = [[str(labels["assumption"]), str(labels["value"]), str(labels["unit"])]]
    appendix_assumptions = (
        ("business_name", assumptions.business_name, ""),
        (
            "restaurant_format",
            formats[assumptions.restaurant_format],
            "",
        ),
        ("location", assumptions.location, ""),
        ("seats", assumptions.seats, ""),
        ("menu_summary", assumptions.menu_summary, ""),
        ("average_check", assumptions.average_check, str(labels["currency"])),
        ("daily_customers", assumptions.daily_customers, ""),
        ("operating_days_per_month", assumptions.operating_days_per_month, ""),
        ("monthly_rent", assumptions.monthly_rent, str(labels["currency"])),
        ("monthly_utilities", assumptions.monthly_utilities, str(labels["currency"])),
        (
            "variable_cost_ratio_percent",
            assumptions.variable_cost_ratio_percent,
            "%",
        ),
        ("other_startup_cost", assumptions.other_startup_cost, str(labels["currency"])),
        ("loan_amount", assumptions.loan_amount, str(labels["currency"])),
        (
            "annual_interest_rate_percent",
            assumptions.annual_interest_rate_percent,
            "%",
        ),
        ("loan_term_months", assumptions.loan_term_months, ""),
        ("target_monthly_profit", assumptions.target_monthly_profit, str(labels["currency"])),
    )
    for key, value, unit in appendix_assumptions:
        if isinstance(value, Decimal):
            rendered_value = f"{value:,.2f}"
        else:
            rendered_value = escape(str(value))
        assumption_rows.append([str(labels["assumptions"][key]), rendered_value, unit])
    story.append(_pdf_table(assumption_rows, [70 * mm, 65 * mm, 25 * mm]))

    story.append(
        Paragraph(
            escape(str(labels["sheets"]["startup_costs"])),
            styles["FinAdvisorHeading"],
        )
    )
    startup_rows = [
        [
            str(labels["columns"]["item"]),
            str(labels["columns"]["quantity"]),
            str(labels["columns"]["unit_cost"]),
            str(labels["columns"]["total_cost"]),
        ]
    ]
    for item in projection.startup_capex.items:
        startup_rows.append(
            [
                escape(item.name),
                str(item.quantity),
                f"{item.unit_cost:,.2f}",
                f"{item.total_cost:,.2f}",
            ]
        )
    startup_rows.append(
        [
            str(labels["assumptions"]["other_startup_cost"]),
            "1",
            f"{projection.other_startup_cost:,.2f}",
            f"{projection.other_startup_cost:,.2f}",
        ]
    )
    startup_rows.append([str(labels["total"]), "", "", f"{projection.total_startup_cost:,.2f}"])
    story.append(_pdf_table(startup_rows, [65 * mm, 25 * mm, 35 * mm, 35 * mm]))

    story.append(
        Paragraph(
            escape(str(labels["columns"]["monthly_payroll"])),
            styles["FinAdvisorHeading"],
        )
    )
    payroll_rows = [
        [
            str(labels["columns"]["role"]),
            str(labels["columns"]["employee_count"]),
            str(labels["columns"]["monthly_salary"]),
            str(labels["columns"]["monthly_payroll"]),
        ]
    ]
    for position in projection.monthly_payroll.positions:
        payroll_rows.append(
            [
                escape(position.role),
                str(position.employee_count),
                f"{position.monthly_salary:,.2f}",
                f"{position.monthly_payroll:,.2f}",
            ]
        )
    payroll_rows.append(
        [
            str(labels["total"]),
            str(projection.monthly_payroll.employee_count),
            "",
            f"{projection.monthly_payroll.monthly_payroll:,.2f}",
        ]
    )
    story.append(_pdf_table(payroll_rows, [45 * mm, 25 * mm, 50 * mm, 40 * mm]))

    loan_schedule = projection.loan_schedule
    story.append(
        Paragraph(
            escape(str(labels["sheets"]["loan_schedule"])),
            styles["FinAdvisorHeading"],
        )
    )
    if loan_schedule is None:
        story.append(
            Paragraph(
                escape(str(labels["status"]["no_loan"])),
                styles["FinAdvisorBody"],
            )
        )
    else:
        loan_rows = [
            [
                str(labels["columns"]["month"]),
                str(labels["columns"]["payment"]),
                str(labels["columns"]["principal"]),
                str(labels["columns"]["interest"]),
                str(labels["columns"]["remaining_balance"]),
            ]
        ]
        for payment in loan_schedule.payments:
            loan_rows.append(
                [
                    str(payment.month),
                    f"{payment.payment:,.2f}",
                    f"{payment.principal:,.2f}",
                    f"{payment.interest:,.2f}",
                    f"{payment.remaining_balance:,.2f}",
                ]
            )
        story.append(_pdf_table(loan_rows, [20 * mm, 35 * mm, 35 * mm, 35 * mm, 35 * mm]))
    story.extend(
        [
            Spacer(1, 8),
            Paragraph(escape(str(labels["disclaimer"])), styles["FinAdvisorBody"]),
        ]
    )
    document.build(story)
    return output.getvalue()


def render_plan_excel(request: PlanDocumentRequest) -> bytes:
    labels = _messages(request.locale)
    assumptions = request.plan.assumptions
    projection = _restaurant_projection(assumptions)
    workbook = Workbook()
    summary = workbook.active
    sheet_names = labels["sheets"]
    summary.title = sheet_names["summary"]
    financials = workbook.create_sheet(sheet_names["financials"])
    financial_sheet_name = str(sheet_names["financials"]).replace("'", "''")
    plan_sheet = workbook.create_sheet(sheet_names["plan"])
    assumption_sheet = workbook.create_sheet(sheet_names["assumptions"])
    startup_sheet = workbook.create_sheet(sheet_names["startup_costs"])
    payroll_sheet = workbook.create_sheet(sheet_names["payroll"])
    monthly_pnl_sheet = workbook.create_sheet(sheet_names["monthly_pnl"])
    loan_sheet = workbook.create_sheet(sheet_names["loan_schedule"])
    break_even_sheet = workbook.create_sheet(sheet_names["break_even"])
    header_fill = PatternFill("solid", fgColor="183B56")
    header_font = Font(color="FFFFFF", bold=True)
    columns = labels["columns"]
    calculated_values = _projection_values(projection)

    financials.append([str(labels["financials"]), labels["value"], labels["unit"]])
    financial_rows: dict[str, int] = {}
    for key in FINANCIAL_KEYS:
        value = calculated_values[key]
        if value is not None:
            financials.append(
                [
                    labels["metrics"][key],
                    value,
                    str(labels["currency"]) if key in CURRENCY_KEYS else "",
                ]
            )
            financial_rows[key] = financials.max_row
            if key in CURRENCY_KEYS:
                financials.cell(financials.max_row, 2).number_format = MONEY_NUMBER_FORMAT
    financials.freeze_panes = "A2"
    financials.column_dimensions["A"].width = 40
    financials.column_dimensions["B"].width = 24
    financials.column_dimensions["C"].width = 14

    summary.append([str(labels["title"]), labels["value"]])
    summary_keys = (
        "monthly_revenue",
        "monthly_operating_profit",
        "annual_operating_profit",
        "total_startup_cost",
        "break_even_revenue",
        "dscr",
    )
    for key in summary_keys:
        if key in financial_rows:
            summary.append(
                [
                    labels["metrics"][key],
                    f"='{financial_sheet_name}'!B{financial_rows[key]}",
                ]
            )
            if key in CURRENCY_KEYS:
                summary.cell(summary.max_row, 2).number_format = MONEY_NUMBER_FORMAT
    summary.freeze_panes = "A2"
    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 24

    assumption_sheet.append([labels["assumption"], labels["value"], labels["unit"]])
    assumption_rows = (
        ("business_name", assumptions.business_name, ""),
        (
            "restaurant_format",
            labels["formats"][assumptions.restaurant_format],
            "",
        ),
        ("location", assumptions.location, ""),
        ("seats", assumptions.seats, ""),
        ("menu_summary", assumptions.menu_summary, ""),
        ("average_check", assumptions.average_check, labels["currency"]),
        ("daily_customers", assumptions.daily_customers, ""),
        ("operating_days_per_month", assumptions.operating_days_per_month, ""),
        ("monthly_rent", assumptions.monthly_rent, labels["currency"]),
        ("monthly_utilities", assumptions.monthly_utilities, labels["currency"]),
        (
            "variable_cost_ratio_percent",
            assumptions.variable_cost_ratio_percent,
            "%",
        ),
        ("other_startup_cost", assumptions.other_startup_cost, labels["currency"]),
        ("loan_amount", assumptions.loan_amount, labels["currency"]),
        (
            "annual_interest_rate_percent",
            assumptions.annual_interest_rate_percent,
            "%",
        ),
        ("loan_term_months", assumptions.loan_term_months, ""),
        ("target_monthly_profit", assumptions.target_monthly_profit, labels["currency"]),
    )
    assumption_sheet_rows: dict[str, int] = {}
    for key, value, unit in assumption_rows:
        assumption_sheet.append(
            [
                labels["assumptions"][key],
                _excel_text(value) if isinstance(value, str) else value,
                unit,
            ]
        )
        assumption_sheet_rows[key] = assumption_sheet.max_row
        if unit == labels["currency"]:
            assumption_sheet.cell(assumption_sheet.max_row, 2).number_format = MONEY_NUMBER_FORMAT
    assumption_sheet.freeze_panes = "A2"
    assumption_sheet.column_dimensions["A"].width = 36
    assumption_sheet.column_dimensions["B"].width = 60
    assumption_sheet.column_dimensions["C"].width = 14

    startup_sheet.append(
        [
            columns["item"],
            columns["quantity"],
            columns["unit_cost"],
            columns["total_cost"],
        ]
    )
    for item in projection.startup_capex.items:
        startup_sheet.append(
            [
                _excel_text(item.name),
                item.quantity,
                item.unit_cost,
                item.total_cost,
            ]
        )
    startup_sheet.append(
        [
            labels["assumptions"]["other_startup_cost"],
            Decimal(1),
            projection.other_startup_cost,
            projection.other_startup_cost,
        ]
    )
    startup_sheet.append([labels["total"], None, None, projection.total_startup_cost])
    for row in range(2, startup_sheet.max_row + 1):
        startup_sheet.cell(row, 3).number_format = MONEY_NUMBER_FORMAT
        startup_sheet.cell(row, 4).number_format = MONEY_NUMBER_FORMAT
    startup_sheet.freeze_panes = "A2"
    startup_sheet.column_dimensions["A"].width = 40
    startup_sheet.column_dimensions["B"].width = 16
    startup_sheet.column_dimensions["C"].width = 24
    startup_sheet.column_dimensions["D"].width = 24

    payroll_sheet.append(
        [
            columns["role"],
            columns["employee_count"],
            columns["monthly_salary"],
            columns["monthly_payroll"],
        ]
    )
    for position in projection.monthly_payroll.positions:
        payroll_sheet.append(
            [
                _excel_text(position.role),
                position.employee_count,
                position.monthly_salary,
                position.monthly_payroll,
            ]
        )
    payroll_sheet.append(
        [
            labels["total"],
            projection.monthly_payroll.employee_count,
            None,
            projection.monthly_payroll.monthly_payroll,
        ]
    )
    for row in range(2, payroll_sheet.max_row + 1):
        payroll_sheet.cell(row, 3).number_format = MONEY_NUMBER_FORMAT
        payroll_sheet.cell(row, 4).number_format = MONEY_NUMBER_FORMAT
    payroll_sheet.append([])
    payroll_contribution_header = payroll_sheet.max_row + 1
    payroll_sheet.append(
        [
            columns["contribution"],
            columns["rate"],
            columns["monthly_amount"],
            columns["verification"],
            columns["source"],
            columns["effective_date"],
        ]
    )
    for contribution in projection.monthly_payroll.employer_contributions:
        status_label = (
            labels["status"]["verified"]
            if contribution.verified
            else (
                labels["status"]["not_supplied"]
                if contribution.rate is None
                else labels["status"]["unverified"]
            )
        )
        payroll_sheet.append(
            [
                _excel_text(contribution.key),
                contribution.rate,
                contribution.monthly_amount,
                status_label,
                _excel_text(contribution.source_url),
                _excel_text(contribution.effective_date),
            ]
        )
        payroll_sheet.cell(payroll_sheet.max_row, 3).number_format = MONEY_NUMBER_FORMAT
    payroll_sheet.freeze_panes = "A2"
    for column, width in {
        "A": 34,
        "B": 18,
        "C": 24,
        "D": 26,
        "E": 52,
        "F": 18,
    }.items():
        payroll_sheet.column_dimensions[column].width = width

    monthly_pnl_sheet.append([labels["monthly_pnl_note"]])
    monthly_pnl_sheet.append([labels["monthly_pnl_debt_note"]])
    monthly_pnl_sheet.append([])
    month_headers = [
        f"{labels['month']} {month}" for month in range(1, MONTHLY_PROJECTION_MONTHS + 1)
    ]
    monthly_pnl_sheet.append([columns["item"], *month_headers])
    financial_sheet_reference = str(sheet_names["financials"]).replace("'", "''")
    assumption_sheet_reference = str(sheet_names["assumptions"]).replace("'", "''")
    pnl_rows = (
        (
            labels["metrics"]["monthly_revenue"],
            financial_sheet_reference,
            financial_rows["monthly_revenue"],
        ),
        (
            labels["metrics"]["monthly_payroll"],
            financial_sheet_reference,
            financial_rows["monthly_payroll"],
        ),
        (
            labels["assumptions"]["monthly_rent"],
            assumption_sheet_reference,
            assumption_sheet_rows["monthly_rent"],
        ),
        (
            labels["assumptions"]["monthly_utilities"],
            assumption_sheet_reference,
            assumption_sheet_rows["monthly_utilities"],
        ),
        (
            labels["metrics"]["monthly_variable_cost"],
            financial_sheet_reference,
            financial_rows["monthly_variable_cost"],
        ),
        (
            labels["metrics"]["monthly_operating_profit"],
            financial_sheet_reference,
            financial_rows["monthly_operating_profit"],
        ),
    )
    for row_label, source_sheet, source_row in pnl_rows:
        monthly_pnl_sheet.append([row_label, *([None] * MONTHLY_PROJECTION_MONTHS)])
        pnl_row = monthly_pnl_sheet.max_row
        for column_index in range(2, MONTHLY_PROJECTION_MONTHS + 2):
            cell = monthly_pnl_sheet.cell(pnl_row, column_index)
            cell.value = f"='{source_sheet}'!B{source_row}"
            cell.number_format = MONEY_NUMBER_FORMAT
    if projection.loan_schedule is not None:
        loan_payments = projection.loan_schedule.payments[:MONTHLY_PROJECTION_MONTHS]
        monthly_pnl_sheet.append(
            [
                labels["metrics"]["loan_first_payment"],
                *([None] * MONTHLY_PROJECTION_MONTHS),
            ]
        )
        pnl_row = monthly_pnl_sheet.max_row
        loan_sheet_reference = str(sheet_names["loan_schedule"]).replace("'", "''")
        for month in range(1, MONTHLY_PROJECTION_MONTHS + 1):
            cell = monthly_pnl_sheet.cell(pnl_row, month + 1)
            cell.value = (
                f"='{loan_sheet_reference}'!B{month + 1}"
                if month <= len(loan_payments)
                else Decimal(0)
            )
            cell.number_format = MONEY_NUMBER_FORMAT
    monthly_pnl_sheet.freeze_panes = "B5"
    monthly_pnl_sheet.column_dimensions["A"].width = 34
    for column_index in range(2, 14):
        monthly_pnl_sheet.column_dimensions[chr(ord("A") + column_index - 1)].width = 18

    if projection.loan_schedule is None:
        loan_sheet.append([labels["status"]["no_loan"]])
    else:
        loan_sheet.append(
            [
                columns["month"],
                columns["payment"],
                columns["principal"],
                columns["interest"],
                columns["remaining_balance"],
            ]
        )
        for payment in projection.loan_schedule.payments:
            loan_sheet.append(
                [
                    payment.month,
                    payment.payment,
                    payment.principal,
                    payment.interest,
                    payment.remaining_balance,
                ]
            )
            for column_index in range(2, 6):
                loan_sheet.cell(
                    loan_sheet.max_row, column_index
                ).number_format = MONEY_NUMBER_FORMAT
        loan_sheet.freeze_panes = "A2"
        for column in ("A", "B", "C", "D", "E"):
            loan_sheet.column_dimensions[column].width = 24

    break_even_sheet.append([labels["assumption"], labels["value"], labels["unit"]])
    break_even_sheet.append(
        [
            labels["metrics"]["monthly_fixed_costs"],
            projection.monthly_fixed_costs,
            labels["currency"],
        ]
    )
    break_even_sheet.append(
        [
            labels["assumptions"]["variable_cost_ratio_percent"],
            assumptions.variable_cost_ratio_percent,
            "%",
        ]
    )
    break_even_sheet.append(
        [
            labels["metrics"]["break_even_revenue"],
            projection.break_even.break_even_revenue,
            labels["currency"],
        ]
    )
    break_even_sheet.append([columns["formula"], _excel_text(projection.break_even.formula), ""])
    break_even_sheet["B2"].number_format = MONEY_NUMBER_FORMAT
    break_even_sheet["B4"].number_format = MONEY_NUMBER_FORMAT
    break_even_sheet.freeze_panes = "A2"
    break_even_sheet.column_dimensions["A"].width = 40
    break_even_sheet.column_dimensions["B"].width = 24
    break_even_sheet.column_dimensions["C"].width = 14

    plan_sheet.append([str(labels["plan"]), ""])
    plan_sheet.append([str(labels["summary"]), _excel_text(request.plan.summary)])
    plan_sheet.append(["", ""])
    for section in request.plan.sections:
        plan_sheet.append([_excel_text(section.title), _excel_text(section.content)])
    plan_sheet.column_dimensions["A"].width = 30
    plan_sheet.column_dimensions["B"].width = 100
    for sheet in (
        summary,
        financials,
        plan_sheet,
        assumption_sheet,
        startup_sheet,
        payroll_sheet,
        loan_sheet,
        break_even_sheet,
    ):
        _style_header(sheet, 1, header_fill, header_font)
    _style_header(monthly_pnl_sheet, 4, header_fill, header_font)
    if projection.monthly_payroll.employer_contributions:
        _style_header(payroll_sheet, payroll_contribution_header, header_fill, header_font)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.post("/{plan_id}/pdf")
def export_plan_pdf(
    request: PlanDocumentRequest,
    user: Annotated[User, Depends(get_current_user)],
    plan_id: str = Path(pattern=r"^[A-Za-z0-9_-]{1,80}$"),
) -> Response:
    require_plan_feature(user, "exports")
    content = render_plan_pdf(request)
    return Response(
        content=content,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="finadvisor-plan-{plan_id}.pdf"'},
    )


@router.post("/{plan_id}/bank-pack")
def export_bank_pack_pdf(
    request: PlanDocumentRequest,
    user: Annotated[User, Depends(get_current_user)],
    plan_id: str = Path(pattern=r"^[A-Za-z0-9_-]{1,80}$"),
) -> Response:
    require_plan_feature(user, "exports")
    content = render_bank_pack_pdf(request)
    return Response(
        content=content,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (f'attachment; filename="finadvisor-bank-pack-{plan_id}.pdf"')
        },
    )


@router.post("/{plan_id}/synergy")
def export_synergy_pdf(
    request: PlanDocumentRequest,
    user: Annotated[User, Depends(get_current_user)],
    plan_id: str = Path(pattern=r"^[A-Za-z0-9_-]{1,80}$"),
) -> Response:
    require_plan_feature(user, "agent_synergy")
    content = render_synergy_pdf(request)
    return Response(
        content=content,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                f'attachment; filename="finadvisor-agent-synergy-{plan_id}.pdf"'
            )
        },
    )


@router.post("/{plan_id}/excel")
def export_plan_excel(
    request: PlanDocumentRequest,
    user: Annotated[User, Depends(get_current_user)],
    plan_id: str = Path(pattern=r"^[A-Za-z0-9_-]{1,80}$"),
) -> Response:
    require_plan_feature(user, "exports")
    content = render_plan_excel(request)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="finadvisor-plan-{plan_id}.xlsx"'},
    )
