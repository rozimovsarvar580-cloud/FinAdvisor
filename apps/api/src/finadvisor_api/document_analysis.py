from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from finadvisor_api.readiness import ReadinessResult, calculate_bank_readiness

MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MAX_DOCUMENT_PAGES = 100
MAX_ARCHIVE_ENTRIES = 2_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_SPREADSHEET_SHEETS = 20
MAX_SPREADSHEET_ROWS = 10_000
MAX_SPREADSHEET_COLUMNS = 100
MAX_SPREADSHEET_CELLS = 100_000


class DocumentExtractionError(ValueError):
    pass


def _office_archive_members(content: bytes) -> set[str]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_ARCHIVE_ENTRIES
                or sum(entry.file_size for entry in entries)
                > MAX_ARCHIVE_UNCOMPRESSED_BYTES
            ):
                raise DocumentExtractionError("analysis.file_too_large")
            return {entry.filename for entry in entries}
    except BadZipFile as error:
        raise DocumentExtractionError("analysis.invalid_document") from error


def _detect_document_type(content: bytes) -> str:
    if content.startswith(b"%PDF-"):
        return "pdf"

    if content.startswith(b"PK"):
        members = _office_archive_members(content)
        if "word/document.xml" in members:
            return "docx"
        if "xl/workbook.xml" in members:
            return "xlsx"
        raise DocumentExtractionError("analysis.unsupported_file_type")

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise DocumentExtractionError("analysis.unsupported_file_type") from error
    if "\x00" in text:
        raise DocumentExtractionError("analysis.unsupported_file_type")
    return "txt"


def _extract_docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend(
            "\t".join(cell.text for cell in row.cells)
            for row in table.rows
        )
    return "\n".join(parts)


def _extract_xlsx_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) > MAX_SPREADSHEET_SHEETS:
            raise DocumentExtractionError("analysis.file_too_large")

        parts: list[str] = []
        cells_read = 0
        for worksheet in workbook.worksheets:
            sheet_parts: list[str] = []
            for row_index, row in enumerate(
                worksheet.iter_rows(
                    max_col=MAX_SPREADSHEET_COLUMNS,
                    values_only=True,
                ),
                start=1,
            ):
                if row_index > MAX_SPREADSHEET_ROWS:
                    raise DocumentExtractionError("analysis.file_too_large")
                cells_read += len(row)
                if cells_read > MAX_SPREADSHEET_CELLS:
                    raise DocumentExtractionError("analysis.file_too_large")
                values = [str(value) for value in row if value is not None]
                if values:
                    if not sheet_parts:
                        sheet_parts.append(worksheet.title)
                    sheet_parts.append("\t".join(values))
            parts.extend(sheet_parts)
        return "\n".join(parts)
    finally:
        workbook.close()


def extract_document_text(filename: str, content: bytes) -> str:
    if len(content) > MAX_DOCUMENT_BYTES:
        raise DocumentExtractionError("analysis.file_too_large")
    if not content:
        raise DocumentExtractionError("analysis.empty_file")

    document_type = _detect_document_type(content)
    try:
        if document_type == "pdf":
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise DocumentExtractionError("analysis.invalid_document")
            if len(reader.pages) > MAX_DOCUMENT_PAGES:
                raise DocumentExtractionError("analysis.too_many_pages")
            extracted_text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )
        elif document_type == "docx":
            extracted_text = _extract_docx_text(content)
        elif document_type == "xlsx":
            extracted_text = _extract_xlsx_text(content)
        else:
            extracted_text = content.decode("utf-8-sig")
    except DocumentExtractionError:
        raise
    except (
        PdfReadError,
        PackageNotFoundError,
        InvalidFileException,
        BadZipFile,
        XMLSyntaxError,
        UnicodeDecodeError,
        KeyError,
        ValueError,
        OSError,
        RuntimeError,
    ) as error:
        raise DocumentExtractionError("analysis.invalid_document") from error

    if not extracted_text.strip():
        raise DocumentExtractionError("analysis.no_text_found")
    return extracted_text


def analyze_business_plan(filename: str, content: bytes) -> tuple[ReadinessResult, int]:
    text = extract_document_text(filename, content)
    return calculate_bank_readiness(text), len(text.split())
