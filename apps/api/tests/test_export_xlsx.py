from decimal import Decimal
from io import BytesIO

import pytest
from finance_engine.operations import CapexItem, StaffPosition
from finance_engine.restaurant import calculate_restaurant_projection
from openpyxl import load_workbook

from finadvisor_api.documents import PlanDocumentRequest, render_plan_excel


@pytest.mark.parametrize("locale", ["en", "ru", "uz"])
def test_workbook_contains_engine_outputs_and_opens_with_openpyxl(
    locale: str,
) -> None:
    assumptions = {
        "locale": locale,
        "business_name": "Test restaurant",
        "restaurant_format": "family",
        "location": "Tashkent",
        "seats": 24,
        "menu_summary": "Local dishes",
        "average_check": "10000",
        "daily_customers": 10,
        "operating_days_per_month": 30,
        "monthly_rent": "100000",
        "monthly_utilities": "50000",
        "variable_cost_ratio_percent": "30",
        "staff": [{"role": "Cook", "employee_count": 1, "monthly_salary": "500000"}],
        "equipment": [{"name": "Oven", "quantity": "2", "unit_cost": "1000000"}],
        "other_startup_cost": "500000",
        "loan_amount": "1200000",
        "annual_interest_rate_percent": "12",
        "loan_term_months": 12,
        "target_monthly_profit": "100000",
    }
    projection = calculate_restaurant_projection(
        average_check=Decimal(assumptions["average_check"]),
        daily_customers=assumptions["daily_customers"],
        operating_days_per_month=assumptions["operating_days_per_month"],
        monthly_rent=Decimal(assumptions["monthly_rent"]),
        monthly_utilities=Decimal(assumptions["monthly_utilities"]),
        variable_cost_ratio_percent=Decimal(assumptions["variable_cost_ratio_percent"]),
        staff=[StaffPosition("Cook", 1, Decimal(500000))],
        equipment=[CapexItem("Oven", Decimal(2), Decimal(1000000))],
        other_startup_cost=Decimal(500000),
        loan_amount=Decimal(1200000),
        annual_interest_rate_percent=Decimal(12),
        loan_term_months=12,
        target_monthly_profit=Decimal(100000),
    )
    calculations = {
        "monthly_revenue": format(projection.monthly_revenue, "f"),
        "monthly_payroll": format(projection.monthly_payroll.monthly_payroll, "f"),
        "startup_capex": format(projection.startup_capex.total_cost, "f"),
        "total_startup_cost": format(projection.total_startup_cost, "f"),
        "monthly_fixed_costs": format(projection.monthly_fixed_costs, "f"),
        "monthly_variable_cost": format(projection.monthly_variable_cost, "f"),
        "monthly_operating_profit": format(projection.monthly_operating_profit, "f"),
        "annual_operating_profit": format(projection.annual_operating_profit, "f"),
        "break_even_revenue": format(projection.break_even.break_even_revenue, "f"),
        "target_monthly_revenue": format(projection.target_revenue.required_revenue, "f"),
        "loan_first_payment": format(projection.loan_schedule.monthly_payment, "f"),
        "loan_total_payment": format(projection.loan_schedule.total_payment, "f"),
        "loan_total_interest": format(projection.loan_schedule.total_interest, "f"),
        "loan_term_months": projection.loan_schedule.term_months,
        "annual_debt_service": format(projection.annual_debt_service, "f"),
        "dscr": format(projection.dscr.dscr, "f"),
    }
    calculations["monthly_revenue"] = "1"
    request = PlanDocumentRequest.model_validate(
        {
            "locale": locale,
            "plan": {
                "summary": "Test plan",
                "sections": [{"title": "Operations", "content": "Daily service."}],
                "markdown": "# Test plan",
                "calculations": calculations,
                "assumptions": assumptions,
            },
        }
    )

    workbook_bytes = render_plan_excel(request)
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)

    expected_sheets = {
        "en": [
            "Summary",
            "Financials",
            "Plan",
            "Assumptions",
            "Startup costs",
            "Payroll",
            "Monthly P&L",
            "Loan schedule",
            "Break-even",
        ],
        "ru": [
            "Сводка",
            "Финансы",
            "Бизнес-план",
            "Допущения",
            "Стартовые затраты",
            "Зарплата",
            "P&L по месяцам",
            "График кредита",
            "Безубыточность",
        ],
        "uz": [
            "Xulosa",
            "Moliya",
            "Reja",
            "Taxminlar",
            "Boshlang‘ich xarajatlar",
            "Ish haqi",
            "Oylik P&L",
            "Kredit jadvali",
            "Zararsizlik",
        ],
    }
    assert workbook.sheetnames == expected_sheets[locale]
    summary_name, financial_name = workbook.sheetnames[:2]
    startup_name, payroll_name, pnl_name, loan_name, break_even_name = workbook.sheetnames[4:]
    assert workbook[summary_name]["B2"].value == f"='{financial_name}'!B2"
    assert workbook[workbook.sheetnames[3]]["B2"].value == "Test restaurant"
    assert Decimal(str(workbook[financial_name]["B2"].value)) == projection.monthly_revenue
    assert Decimal(str(workbook[startup_name]["D4"].value)) == Decimal("2500000.00")
    assert Decimal(str(workbook[payroll_name]["D3"].value)) == Decimal("500000.00")
    assert workbook[pnl_name]["B5"].value == f"='{financial_name}'!B2"
    assert workbook[pnl_name]["M5"].value == f"='{financial_name}'!B2"
    assert Decimal(str(workbook[loan_name]["B2"].value)) == (
        projection.loan_schedule.payments[0].payment
    )
    assert Decimal(str(workbook[break_even_name]["B4"].value)).quantize(
        Decimal("0.01")
    ) == projection.break_even.break_even_revenue.quantize(Decimal("0.01"))
    assert workbook[payroll_name]["D3"].number_format == "#,##0.00"
