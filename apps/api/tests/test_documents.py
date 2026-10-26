from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from finadvisor_api.auth import get_current_user
from finadvisor_api.documents import (
    DOCUMENT_MESSAGES,
    PlanDocumentRequest,
    _excel_text,
    render_plan_excel,
    render_plan_pdf,
)
from finadvisor_api.main import app
from finadvisor_api.models import User

client = TestClient(app)

plan_response = {
    "assumptions": {
        "locale": "uz",
        "business_name": "Oshxona",
        "restaurant_format": "family",
        "location": "Tashkent",
        "seats": 40,
        "menu_summary": "Local dishes",
        "average_check": "100000",
        "daily_customers": 50,
        "operating_days_per_month": 30,
        "monthly_rent": "10000000",
        "monthly_utilities": "1000000",
        "variable_cost_ratio_percent": "30",
        "staff": [
            {"role": "Cook", "employee_count": 10, "monthly_salary": "3000000"},
            {"role": "Server", "employee_count": 20, "monthly_salary": "2000000"},
        ],
        "equipment": [{"name": "Kitchen equipment", "quantity": "1", "unit_cost": "50000000"}],
        "other_startup_cost": "10000000",
        "loan_amount": "0",
        "annual_interest_rate_percent": "0",
        "loan_term_months": 0,
        "target_monthly_profit": "10000000",
    },
    "summary": "A restaurant serving local cuisine.",
    "sections": [
        {"title": "Market overview", "content": "The restaurant serves local dishes."},
        {"title": "Operations", "content": "The restaurant will operate daily."},
        {"title": "Staffing", "content": "The plan includes cooks and servers."},
        {"title": "Risks", "content": "Risks require ongoing review."},
        {"title": "Financial overview", "content": "Engine-calculated values follow."},
    ],
    "markdown": (
        "# A restaurant serving local cuisine.\n\n"
        "## Operations\n\nThe restaurant will operate daily."
    ),
    "calculations": {
        "monthly_revenue": "150000000.00",
        "monthly_payroll": "70000000.00",
        "startup_capex": "50000000.00",
        "total_startup_cost": "60000000.00",
        "monthly_fixed_costs": "81000000.00",
        "monthly_variable_cost": "45000000.00",
        "monthly_operating_profit": "24000000.00",
        "annual_operating_profit": "288000000.00",
        "break_even_revenue": "115714285.71",
        "target_monthly_revenue": "130000000.00",
        "loan_first_payment": None,
        "loan_total_payment": None,
        "loan_total_interest": None,
        "loan_term_months": None,
        "annual_debt_service": None,
        "dscr": None,
    },
}
document_request = {"locale": "uz", "plan": plan_response}


def authorize_export_plan(monkeypatch: pytest.MonkeyPatch, plan: str = "pro") -> User:
    user = User(
        id="document-export-user",
        email="document-export@example.com",
        name="Document Export",
        role="tadbirkor",
        subscription_plan=plan,
    )
    monkeypatch.setitem(app.dependency_overrides, get_current_user, lambda: user)
    return user


@pytest.mark.parametrize("locale", ["uz", "ru", "en"])
def test_render_plan_pdf_returns_a_pdf_with_localized_document_content(locale: str) -> None:
    request = PlanDocumentRequest.model_validate({**document_request, "locale": locale})

    result = render_plan_pdf(request)
    reader = PdfReader(BytesIO(result))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)

    assert result.startswith(b"%PDF-")
    assert len(result) > 1_000
    assert len(reader.pages) >= 3
    assert "Oshxona" in text
    assert "The restaurant serves local dishes." in text
    assert "The restaurant will operate daily." in text
    assert "The plan includes cooks and servers." in text
    assert "Risks require ongoing review." in text
    assert "150,000,000.00 UZS" in text
    assert "60,000,000.00" in text

    localized_headings = DOCUMENT_MESSAGES[locale]
    assert all(
        localized_headings[key] in text
        for key in ("market", "operations", "staffing", "risks", "appendix")
    ), text


def test_pdf_uses_engine_results_instead_of_submitted_calculations() -> None:
    mismatched_plan = {
        **plan_response,
        "calculations": {**plan_response["calculations"], "monthly_revenue": "1"},
    }
    request = PlanDocumentRequest.model_validate({"locale": "en", "plan": mismatched_plan})

    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(render_plan_pdf(request))).pages
    )

    assert "150,000,000.00 UZS" in text
    assert "1.00 UZS" not in text


def test_render_plan_excel_has_multiple_sheets_and_engine_linked_formulas() -> None:
    request = PlanDocumentRequest.model_validate(document_request)

    result = render_plan_excel(request)
    workbook = load_workbook(BytesIO(result), data_only=False)

    assert workbook.sheetnames == [
        "Xulosa",
        "Moliya",
        "Reja",
        "Taxminlar",
        "Boshlang‘ich xarajatlar",
        "Ish haqi",
        "Oylik P&L",
        "Kredit jadvali",
        "Zararsizlik",
    ]
    assert workbook["Xulosa"]["B2"].value == "='Moliya'!B2"
    assert Decimal(str(workbook["Moliya"]["B2"].value)) == Decimal("150000000.00")
    assert workbook["Reja"]["B2"].value == plan_response["summary"]
    assert workbook["Taxminlar"]["B2"].value == "Oshxona"
    assert Decimal(str(workbook["Boshlang‘ich xarajatlar"]["D4"].value)) == Decimal("60000000.00")
    assert Decimal(str(workbook["Ish haqi"]["D4"].value)) == Decimal("70000000.00")
    monthly_pnl = workbook["Oylik P&L"]
    assert monthly_pnl["A4"].value == "Mahsulot"
    assert [monthly_pnl.cell(row=5, column=column).value for column in range(2, 14)] == [
        "='Moliya'!B2"
    ] * 12
    assert Decimal(str(workbook["Moliya"]["B2"].value)) == Decimal("150000000.00")
    assert Decimal(str(workbook["Zararsizlik"]["B4"].value)).quantize(Decimal("0.01")) == Decimal(
        "115714285.71"
    )


def test_excel_text_neutralizes_formula_like_plan_content() -> None:
    assert _excel_text('=HYPERLINK("https://example.invalid")').startswith("'=")
    assert _excel_text('  =HYPERLINK("https://example.invalid")').startswith("'  =")
    assert _excel_text("\t=HYPERLINK").startswith("'\t=")
    assert _excel_text("ordinary text") == "ordinary text"


def test_document_request_rejects_non_decimal_financial_values() -> None:
    payload = {
        **document_request,
        "plan": {
            **plan_response,
            "calculations": {
                **plan_response["calculations"],
                "monthly_revenue": '=HYPERLINK("https://example.invalid")',
            },
        },
    }

    with pytest.raises(ValueError, match="monthly_revenue must be a decimal value"):
        PlanDocumentRequest.model_validate(payload)


def test_pdf_export_endpoint_sets_type_and_safe_attachment_name(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    authorize_export_plan(monkeypatch)
    response = client.post(
        "/documents/plan_123/pdf",
        json=document_request,
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        'attachment; filename="finadvisor-plan-plan_123.pdf"'
    )
    assert response.content.startswith(b"%PDF-")


def test_document_export_requires_authentication() -> None:
    response = client.post("/documents/plan_123/pdf", json=document_request)

    assert response.status_code == 401


@pytest.mark.parametrize("format", ["pdf", "excel", "bank-pack"])
def test_free_plan_cannot_export_documents(format: str, monkeypatch: pytest.MonkeyPatch) -> None:
    authorize_export_plan(monkeypatch, plan="free")

    response = client.post(f"/documents/plan_123/{format}", json=document_request)

    assert response.status_code == 403
    assert response.json()["detail"] == {
        "code": "plans.feature_not_included",
        "feature": "exports",
        "upgrade_hint": "pro",
    }


def test_bank_pack_export_includes_loan_schedule_collateral_and_readiness(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    authorize_export_plan(monkeypatch)
    assumptions = {
        **plan_response["assumptions"],
        "locale": "en",
        "loan_amount": "1200000",
        "annual_interest_rate_percent": "12",
        "loan_term_months": 12,
    }
    request = {
        "locale": "en",
        "plan": {**plan_response, "assumptions": assumptions},
    }

    response = client.post("/documents/plan_123/bank-pack", json=request)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        'attachment; filename="finadvisor-bank-pack-plan_123.pdf"'
    )
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages
    )
    assert "Loan request" in text
    assert "1,200,000.00 UZS" in text
    assert "Repayment schedule" in text
    assert "Collateral details were not supplied" in text
    assert "Bank readiness evidence score" in text
    assert "not a bank decision" in text


def test_synergy_export_includes_operating_strategy_and_engine_financials(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    authorize_export_plan(monkeypatch, plan="business")
    request = {
        **document_request,
        "locale": "en",
        "plan": {
            **plan_response,
            "assumptions": {**plan_response["assumptions"], "locale": "en"},
        },
    }

    response = client.post("/documents/plan_123/synergy", json=request)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        'attachment; filename="finadvisor-agent-synergy-plan_123.pdf"'
    )
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(response.content)).pages
    )
    assert "Agent-Synergy operating strategy" in text
    assert "The restaurant will operate daily." in text
    assert "Staffing and responsibilities" in text
    assert "Operating financial signals" in text
    assert "150,000,000.00 UZS" in text


@pytest.mark.parametrize("locale", ["uz", "ru", "en"])
def test_bank_pack_and_synergy_headings_are_localized(
    locale: str,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    user = authorize_export_plan(monkeypatch)
    request = {
        "locale": locale,
        "plan": {
            **plan_response,
            "assumptions": {**plan_response["assumptions"], "locale": locale},
        },
    }
    messages = DOCUMENT_MESSAGES[locale]

    bank_pack = client.post("/documents/plan_123/bank-pack", json=request)
    user.subscription_plan = "business"
    synergy = client.post("/documents/plan_123/synergy", json=request)
    bank_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(bank_pack.content)).pages
    )
    synergy_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(synergy.content)).pages
    )

    assert bank_pack.status_code == 200
    assert synergy.status_code == 200
    assert messages["bank_pack_title"] in bank_text
    assert messages["collateral"] in bank_text
    assert messages["bank_readiness"] in bank_text
    assert messages["synergy_title"] in synergy_text
    assert messages["operating_plan"] in synergy_text
    assert messages["operating_financials"] in synergy_text


def test_excel_export_endpoint_returns_workbook_attachment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    authorize_export_plan(monkeypatch)
    response = client.post(
        "/documents/plan_123/excel",
        json=document_request,
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"].endswith('.xlsx"')
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert "Oylik P&L" in workbook.sheetnames
    assert "Kredit jadvali" in workbook.sheetnames


def test_export_rejects_unsafe_plan_id() -> None:
    response = client.post(
        "/documents/../pdf",
        json=document_request,
    )

    assert response.status_code in {404, 422}
