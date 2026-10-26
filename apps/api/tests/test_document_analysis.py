import json
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from finadvisor_api import analysis
from finadvisor_api.ai_client import AIClientError
from finadvisor_api.auth import get_current_user
from finadvisor_api.database import Base, get_db
from finadvisor_api.document_analysis import (
    MAX_DOCUMENT_BYTES,
    MAX_DOCUMENT_PAGES,
    DocumentExtractionError,
    analyze_business_plan,
    extract_document_text,
)
from finadvisor_api.main import app
from finadvisor_api.models import User


@pytest.fixture
def analysis_client(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    session_factory = sessionmaker(
        bind=engine,
        autoflush=False,
        autocommit=False,
        expire_on_commit=False,
    )
    user = User(
        id="analysis-test-user",
        email="analysis-test@example.com",
        name="Analysis Test",
        role="tadbirkor",
        subscription_plan="pro",
    )
    with session_factory() as db:
        db.add(user)
        db.commit()

    def override_get_db():
        with session_factory() as session:
            yield session

    monkeypatch.setitem(app.dependency_overrides, get_db, override_get_db)
    monkeypatch.setitem(app.dependency_overrides, get_current_user, lambda: user)
    with TestClient(app) as test_client:
        yield test_client
    Base.metadata.drop_all(bind=engine)
    engine.dispose()


def _docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Business plan"
    worksheet.append(["Restaurant", "Revenue"])
    worksheet.append(["Cafe", 150000000])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def _pdf_bytes() -> bytes:
    content = BytesIO()
    document = canvas.Canvas(content)
    document.drawString(72, 720, "Restaurant business plan")
    document.save()
    return content.getvalue()


def test_extracts_docx_text_and_scores_bank_readiness() -> None:
    content = _docx_bytes(
        "Restaurant business overview. Market and customers. "
        "Staff team. Financial revenue and profit. Loan funding. "
        "Risk mitigation, collateral, legal license."
    )
    result, word_count = analyze_business_plan("plan.docx", content)

    assert result.score == 100
    assert word_count > 0


def test_detects_pdf_docx_xlsx_and_txt_from_content() -> None:
    pdf_text = extract_document_text("plan.txt", _pdf_bytes())
    docx_text = extract_document_text(
        "plan.pdf",
        _docx_bytes("Restaurant DOCX content"),
    )
    xlsx_text = extract_document_text("plan.docx", _xlsx_bytes())
    txt_text = extract_document_text("plan.xlsx", b"Restaurant TXT content")

    assert "Restaurant business plan" in pdf_text
    assert "Restaurant DOCX content" in docx_text
    assert "Business plan" in xlsx_text
    assert "150000000" in xlsx_text
    assert txt_text == "Restaurant TXT content"


@pytest.mark.parametrize(
    ("filename", "content", "expected_error"),
    [
        ("plan.bin", b"\x00\xff", "analysis.unsupported_file_type"),
        ("plan.pdf", b"", "analysis.empty_file"),
        ("plan.docx", b"PK\x03\x04invalid", "analysis.invalid_document"),
    ],
)
def test_document_extraction_rejects_invalid_uploads(
    filename: str, content: bytes, expected_error: str
) -> None:
    with pytest.raises(DocumentExtractionError, match=expected_error):
        extract_document_text(filename, content)


def test_document_extraction_enforces_file_size_limit() -> None:
    with pytest.raises(
        DocumentExtractionError,
        match="analysis.file_too_large",
    ):
        extract_document_text("plan.txt", b"x" * (MAX_DOCUMENT_BYTES + 1))


def test_pdf_extraction_enforces_page_limit() -> None:
    writer = PdfWriter()
    for _ in range(MAX_DOCUMENT_PAGES + 1):
        writer.add_blank_page(width=72, height=72)
    content = BytesIO()
    writer.write(content)

    with pytest.raises(DocumentExtractionError, match="analysis.too_many_pages"):
        extract_document_text("plan.pdf", content.getvalue())


def test_analyze_endpoint_accepts_docx_upload(
    analysis_client: TestClient,
) -> None:
    content = _docx_bytes(
        "Restaurant market customers team staff financial revenue profit loan funding "
        "risk mitigation collateral legal license"
    )

    response = analysis_client.post(
        "/plans/analyze",
        files={
            "file": (
                "plan.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["score"] == 100
    assert len(response.json()["criteria"]) == 7


def test_analyze_endpoint_returns_criterion_scores_and_grounded_rewrites(
    monkeypatch: pytest.MonkeyPatch,
    analysis_client: TestClient,
) -> None:
    captured_context: dict[str, object] = {}

    async def fake_generate(prompt: str, context: dict[str, object]) -> str:
        captured_context.update(context)
        missing = context["missing_criteria"]
        assert isinstance(missing, list)
        return json.dumps(
            {
                "improvements": [
                    {
                        "key": key,
                        "fix": "Add supporting information from the business.",
                        "rewritten_section": "[Add verified details here.]",
                    }
                    for key in missing
                ]
            }
        )

    monkeypatch.setattr(analysis, "generate", fake_generate)

    response = analysis_client.post(
        "/plans/analyze",
        data={"locale": "uz"},
        files={"file": ("plan.txt", b"Restaurant market analysis.", "text/plain")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["score"] == 20
    assert body["criteria"][0] == {
        "key": "market_analysis",
        "score": 20,
        "max_score": 20,
        "present": True,
    }
    assert {item["key"] for item in body["improvements"]} == set(
        body["missing_criteria"]
    )
    assert body["improvements"][0]["rewritten_section"] == "[Add verified details here.]"
    assert captured_context["locale"] == "uz"
    assert captured_context["document_text"] == "Restaurant market analysis."


def test_analyze_endpoint_rejects_numeric_ai_rewrites(
    monkeypatch: pytest.MonkeyPatch,
    analysis_client: TestClient,
) -> None:
    async def fake_generate(prompt: str, context: dict[str, object]) -> str:
        missing = context["missing_criteria"]
        assert isinstance(missing, list)
        return json.dumps(
            {
                "improvements": [
                    {
                        "key": key,
                        "fix": "Add details from the business.",
                        "rewritten_section": "Revenue increased by 2026.",
                    }
                    for key in missing
                ]
            }
        )

    monkeypatch.setattr(analysis, "generate", fake_generate)

    response = analysis_client.post(
        "/plans/analyze",
        files={"file": ("plan.txt", b"Restaurant market analysis.", "text/plain")},
    )

    assert response.status_code == 502
    assert response.json()["detail"] == "analysis.invalid_ai_response"


def test_analyze_endpoint_surfaces_ai_provider_failure(
    monkeypatch: pytest.MonkeyPatch,
    analysis_client: TestClient,
) -> None:
    async def unavailable(prompt: str, context: dict[str, object]) -> str:
        raise AIClientError("provider unavailable")

    monkeypatch.setattr(analysis, "generate", unavailable)

    response = analysis_client.post(
        "/plans/analyze",
        files={"file": ("plan.txt", b"Restaurant market analysis.", "text/plain")},
    )

    assert response.status_code == 503
    assert response.json()["detail"] == "analysis.ai_unavailable"
