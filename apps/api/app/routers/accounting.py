import csv
import io

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi import Depends

from ..auth_dependencies import require_role

router = APIRouter(prefix="/accounting", tags=["accounting"])


@router.get("/daily")
async def daily_entries(_: str = Depends(require_role("accountant", "entrepreneur", "admin"))) -> dict[str, object]:
    return {
        "items": [
            {"id": "entry-001", "date": "2026-09-21", "category": "sales", "amount": "8200000", "currency": "UZS"},
            {"id": "entry-002", "date": "2026-09-21", "category": "supplies", "amount": "1650000", "currency": "UZS"},
        ]
    }


@router.get("/inventory")
async def inventory(_: str = Depends(require_role("accountant", "entrepreneur", "admin"))) -> dict[str, object]:
    return {
        "items": [
            {"sku": "rice-001", "name": "Rice", "quantity": "120.00", "unit": "kg", "reorder_level": "50.00"},
            {"sku": "oil-001", "name": "Cooking oil", "quantity": "24.00", "unit": "l", "reorder_level": "30.00"},
        ]
    }


@router.get("/payroll")
async def payroll(_: str = Depends(require_role("accountant", "entrepreneur", "admin"))) -> dict[str, object]:
    return {
        "period": "2026-09",
        "currency": "UZS",
        "employees": [
            {"id": "emp-001", "role": "Chef", "gross_salary": "8500000", "status": "pending"},
            {"id": "emp-002", "role": "Waiter", "gross_salary": "4200000", "status": "pending"},
        ],
    }


@router.get("/tax")
async def tax_summary(_: str = Depends(require_role("accountant", "entrepreneur", "admin"))) -> dict[str, object]:
    return {
        "period": "2026-09",
        "currency": "UZS",
        "regime": "turnover_tax",
        "estimated_due": "12400000",
        "source": "tax_rules_2026.yaml",
        "review_status": "pending_accountant_review",
    }


@router.get("/loans")
async def loans(_: str = Depends(require_role("accountant", "entrepreneur", "admin"))) -> dict[str, object]:
    return {
        "items": [
            {"id": "loan-001", "provider": "Demo Bank", "principal": "500000000", "currency": "UZS", "next_payment": "18500000"},
        ]
    }


@router.post("/import")
async def import_statement(
    file: UploadFile | None = File(default=None),
    _: str = Depends(require_role("accountant", "entrepreneur", "admin")),
) -> dict[str, object]:
    if file is None:
        return {
            "status": "accepted",
            "import_id": "import-001",
            "rows_received": 0,
            "message": "Upload a CSV or XLSX statement to begin import.",
        }
    allowed = {".csv", ".xlsx"}
    suffix = "." + file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if suffix not in allowed:
        raise HTTPException(
            status_code=415,
            detail={"code": "unsupported_file", "message": "Only CSV and XLSX files are supported", "details": {}},
        )
    content = await file.read()
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        row_count = max(0, len(rows) - 1)
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        row_count = max(0, worksheet.max_row - 1)
    return {
        "status": "accepted",
        "import_id": "import-001",
        "rows_received": row_count,
        "filename": file.filename,
        "message": "Statement validated and queued for import.",
    }
